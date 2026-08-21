"""Document text extraction with page-level provenance."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from . import persian

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".markdown", ".htm", ".html"}


@dataclass
class Page:
    number: int
    text: str


@dataclass
class Outline:
    level: int
    title: str
    page: int


@dataclass
class ExtractedDoc:
    title: str
    pages: list[Page]
    outline: list[Outline] = field(default_factory=list)
    scanned_pages: list[int] = field(default_factory=list)
    ocr_pages: list[int] = field(default_factory=list)
    lam_alef_broken: bool = False

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)

    @property
    def full_text(self) -> str:
        return "\n".join(p.text for p in self.pages)


_MISPLACED_DOT_RE = re.compile(r"(?<=[؀-ۿ])\s*([.!؟?])\s*(?=[؀-ۿ])")


def _repair_rtl_punctuation(line: str) -> str:
    """Bidi rendering drops sentence-final punctuation into the middle of RTL lines."""
    fixed, n = _MISPLACED_DOT_RE.subn(" ", line)
    if not n:
        return line
    fixed = fixed.rstrip()
    if fixed and fixed[-1] not in ".!؟?:؛":
        fixed += "."
    return fixed


def _page_lines(page) -> list[str]:
    """Rebuild visual lines from PyMuPDF spans.

    PyMuPDF emits each right-to-left run of a line as its own `line` entry. Joining
    those with newlines splits Persian words apart, so fragments sharing a baseline
    are merged back into one line with spaces.
    """
    try:
        data = page.get_text("dict")
    except Exception:
        return (page.get_text("text") or "").split("\n")

    out: list[str] = []
    for block in data.get("blocks", []):
        rows: list[tuple[float, float, list[str]]] = []
        for ln in block.get("lines", []):
            text = "".join(s.get("text", "") for s in ln.get("spans", []))
            if not text.strip():
                continue
            y0, y1 = ln["bbox"][1], ln["bbox"][3]
            mid = (y0 + y1) / 2
            height = max(y1 - y0, 1.0)
            for row in rows:
                if abs(row[0] - mid) <= max(row[1], height) * 0.6:
                    row[2].append(text)
                    break
            else:
                rows.append((mid, height, [text]))
        for _, _, parts in rows:
            merged = " ".join(p.strip() for p in parts if p.strip())
            if merged.strip():
                out.append(merged)
        out.append("")
    return out


def _clean_page(text_or_lines) -> str:
    lines = text_or_lines if isinstance(text_or_lines, list) else text_or_lines.split("\n")
    cleaned: list[str] = []
    for line in lines:
        line = persian.normalize(persian.fix_visual_order(line))
        if not line:
            continue
        cleaned.append(_repair_rtl_punctuation(line))
    return "\n".join(cleaned)


def _strip_repeated_headers(pages: list[Page]) -> None:
    """Drop header/footer lines that repeat across most pages."""
    if len(pages) < 5:
        return
    from collections import Counter

    first, last = Counter(), Counter()
    for p in pages:
        lines = [ln for ln in p.text.split("\n") if ln]
        if not lines:
            continue
        first[lines[0]] += 1
        last[lines[-1]] += 1
    threshold = max(3, int(len(pages) * 0.4))
    drop = {t for t, c in first.items() if c >= threshold and len(t) < 120}
    drop |= {t for t, c in last.items() if c >= threshold and len(t) < 120}
    if not drop:
        return
    for p in pages:
        lines = [ln for ln in p.text.split("\n") if ln]
        while lines and lines[0] in drop:
            lines.pop(0)
        while lines and lines[-1] in drop:
            lines.pop()
        p.text = "\n".join(lines)


def _ocr_page(page, lang: str) -> str:
    try:
        import pytesseract
        from PIL import Image
        import io
    except Exception:
        return ""
    try:
        pix = page.get_pixmap(dpi=300)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(img, lang=lang)
    except Exception:
        return ""


def extract_pdf(path: Path, ocr: bool = True, ocr_lang: str = "fas+eng") -> ExtractedDoc:
    import pymupdf

    doc = pymupdf.open(path)
    pages: list[Page] = []
    scanned: list[int] = []
    ocred: list[int] = []
    for i, page in enumerate(doc, start=1):
        lines = _page_lines(page)
        if sum(len(ln.strip()) for ln in lines) < 40:
            scanned.append(i)
            if ocr:
                text = _ocr_page(page, ocr_lang)
                if len(text.strip()) >= 40:
                    ocred.append(i)
                    lines = text.split("\n")
        pages.append(Page(number=i, text=_clean_page(lines)))

    outline = []
    try:
        for lvl, title, pno in doc.get_toc(simple=True):
            if pno and pno > 0:
                outline.append(Outline(level=lvl, title=persian.normalize(title), page=pno))
    except Exception:
        pass

    title = ""
    try:
        title = persian.normalize((doc.metadata or {}).get("title") or "")
    except Exception:
        pass
    doc.close()

    _strip_repeated_headers(pages)
    if not title or len(title) < 3:
        title = _guess_title(pages, path)
    return ExtractedDoc(
        title=title,
        pages=pages,
        outline=outline,
        scanned_pages=[p for p in scanned if p not in ocred],
        ocr_pages=ocred,
    )


def _guess_title(pages: list[Page], path: Path) -> str:
    for p in pages[:3]:
        for line in p.text.split("\n"):
            if 6 <= len(line) <= 120:
                return line
    return path.stem


def extract_docx(path: Path) -> ExtractedDoc:
    import docx

    d = docx.Document(str(path))
    blocks: list[str] = []
    outline: list[Outline] = []
    for para in d.paragraphs:
        txt = persian.normalize(para.text)
        if not txt:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading"):
            m = re.search(r"(\d+)", style)
            outline.append(Outline(level=int(m.group(1)) if m else 1, title=txt, page=1))
            blocks.append("\n" + txt)
        else:
            blocks.append(txt)
    for table in d.tables:
        for row in table.rows:
            cells = [persian.normalize(c.text) for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                blocks.append(line)

    full = "\n".join(blocks)
    pages = _paginate(full)
    title = outline[0].title if outline else _guess_title(pages, path)
    return ExtractedDoc(title=title, pages=pages, outline=outline)


def extract_text(path: Path) -> ExtractedDoc:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    if path.suffix.lower() in (".htm", ".html"):
        raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
        raw = re.sub(r"(?s)<[^>]+>", "\n", raw)
        raw = re.sub(r"&nbsp;?", " ", raw)
    text = persian.normalize(raw)
    outline = [
        Outline(level=len(m.group(1)), title=persian.normalize(m.group(2)).strip(), page=1)
        for m in re.finditer(r"^(#{1,6})\s+(.+)$", text, flags=re.M)
    ]
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.M)
    pages = _paginate(text)
    title = outline[0].title if outline else _guess_title(pages, path)
    return ExtractedDoc(title=title, pages=pages, outline=outline)


def _paginate(text: str, chars_per_page: int = 2500) -> list[Page]:
    paras = [p for p in text.split("\n") if p.strip()]
    pages: list[Page] = []
    buf: list[str] = []
    size = 0
    for p in paras:
        buf.append(p)
        size += len(p)
        if size >= chars_per_page:
            pages.append(Page(number=len(pages) + 1, text="\n".join(buf)))
            buf, size = [], 0
    if buf:
        pages.append(Page(number=len(pages) + 1, text="\n".join(buf)))
    return pages or [Page(number=1, text="")]


def extract(path: Path, ocr: bool = True, ocr_lang: str = "fas+eng") -> ExtractedDoc:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        doc = extract_pdf(path, ocr=ocr, ocr_lang=ocr_lang)
    elif suffix == ".docx":
        doc = extract_docx(path)
    elif suffix in SUPPORTED:
        doc = extract_text(path)
    else:
        raise ValueError(f"unsupported file type: {suffix}")
    doc.lam_alef_broken = persian.lam_alef_broken(doc.full_text)
    return doc
